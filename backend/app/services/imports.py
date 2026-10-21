"""Existing-inventory migration: upload → map columns → validate/preview → resolve duplicates → approve.

Opening stock is posted only on approval, by a different user than the uploader, using
deterministic transaction IDs so a retried approval can never double-post stock.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import ImportBatch, ImportRow, Item, ItemCategory, ItemType, Location, MovementType, Uom, User
from .common import BusinessError, audit, next_number, normalize_name
from .items import create_item, find_duplicates
from .stock import MovementInput, post_movement

FIELDS: dict[str, dict] = {
    "item_code": {"label": "Existing item code", "aliases": ["code", "item code", "item no", "part no", "part number", "sku"]},
    "name": {"label": "Item name", "required": True, "aliases": ["item name", "description of goods", "material", "item"]},
    "description": {"label": "Description", "aliases": ["details", "remarks"]},
    "item_type": {"label": "Item type", "required": True, "aliases": ["type", "stream", "material type"]},
    "category": {"label": "Category", "aliases": ["group", "item group"]},
    "uom": {"label": "Unit", "required": True, "aliases": ["unit", "uom", "units", "unit of measure"]},
    "manufacturer": {"label": "Manufacturer", "aliases": ["make", "brand", "mfr"]},
    "make_model": {"label": "Make / model", "aliases": ["model", "make model", "make/model"]},
    "specifications": {"label": "Specifications", "aliases": ["spec", "specs", "specification"]},
    "min_stock": {"label": "Minimum stock", "aliases": ["min", "min stock", "minimum"]},
    "reorder_level": {"label": "Reorder level", "aliases": ["reorder", "rol", "reorder qty"]},
    "max_stock": {"label": "Maximum stock", "aliases": ["max", "max stock", "maximum"]},
    "hsn_code": {"label": "HSN/SAC", "aliases": ["hsn", "sac", "hsn code"]},
    "gst_rate": {"label": "GST %", "aliases": ["gst", "gst rate", "tax rate"]},
    "location": {"label": "Location code", "aliases": ["bin", "location", "rack bin", "store location"]},
    "opening_qty": {"label": "Opening quantity", "aliases": ["qty", "quantity", "stock", "balance", "closing stock", "opening stock"]},
    "unit_cost": {"label": "Unit cost", "aliases": ["rate", "cost", "price", "unit price", "value per unit"]},
}

TYPE_ALIASES = {
    "raw material": ItemType.RAW_MATERIAL, "rm": ItemType.RAW_MATERIAL,
    "purchased item": ItemType.PURCHASED_ITEM, "purchased": ItemType.PURCHASED_ITEM, "bought out": ItemType.PURCHASED_ITEM, "pi": ItemType.PURCHASED_ITEM,
    "daily consumable": ItemType.DAILY_CONSUMABLE, "consumable": ItemType.DAILY_CONSUMABLE, "dc": ItemType.DAILY_CONSUMABLE,
    "processed spare": ItemType.PROCESSED_SPARE, "spare": ItemType.PROCESSED_SPARE, "processed": ItemType.PROCESSED_SPARE, "ps": ItemType.PROCESSED_SPARE,
    "finished good": ItemType.FINISHED_GOOD, "fg": ItemType.FINISHED_GOOD, "machine": ItemType.FINISHED_GOOD,
    "service": ItemType.SERVICE,
}


def parse_file(filename: str, content: bytes) -> tuple[list[str], list[dict]]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = [r for r in reader if any(c.strip() for c in r)]
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)]
    else:
        raise BusinessError("Only CSV and Excel (.xlsx) files are supported")
    if len(rows) < 2:
        raise BusinessError("The file needs a header row and at least one data row")
    headers = [h.strip() or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    records = [{headers[i]: (r[i].strip() if i < len(r) else "") for i in range(len(headers))} for r in rows[1:]]
    if len(records) > 20000:
        raise BusinessError("Files are limited to 20,000 rows per batch")
    return headers, records


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    norm_headers = {normalize_name(h): h for h in headers}
    for field, meta in FIELDS.items():
        for candidate in [field.replace("_", " "), meta["label"].lower(), *meta["aliases"]]:
            hit = norm_headers.get(normalize_name(candidate))
            if hit and hit not in mapping.values():
                mapping[field] = hit
                break
    return mapping


def create_batch(db: Session, user: User, filename: str, content: bytes) -> ImportBatch:
    headers, records = parse_file(filename, content)
    batch = ImportBatch(number=next_number(db, "import_batch", "IMP-", 5), filename=filename, headers=headers,
                        mapping=suggest_mapping(headers), uploaded_by_id=user.id)
    db.add(batch)
    db.flush()
    for i, rec in enumerate(records, start=2):
        db.add(ImportRow(batch_id=batch.id, row_no=i, raw=rec))
    db.flush()
    db.refresh(batch)
    audit(db, user, "import.upload", "import_batch", batch.id, {"filename": filename, "rows": len(records)})
    return batch


def _dec(value: str, field: str, errors: list[str]) -> Decimal | None:
    if value in ("", None):
        return None
    try:
        d = Decimal(str(value).replace(",", ""))
    except InvalidOperation:
        errors.append(f"{FIELDS[field]['label']}: '{value}' is not a number")
        return None
    if d < 0:
        errors.append(f"{FIELDS[field]['label']} cannot be negative")
    return d


def validate_batch(db: Session, user: User, batch: ImportBatch, mapping: dict[str, str]) -> ImportBatch:
    if batch.status not in ("UPLOADED", "VALIDATED"):
        raise BusinessError(f"Batch is {batch.status.lower()} and can no longer be changed", status_code=409, code="state")
    unknown = [f for f in mapping if f not in FIELDS]
    if unknown:
        raise BusinessError(f"Unknown fields in mapping: {', '.join(unknown)}")
    missing = [FIELDS[f]["label"] for f, m in FIELDS.items() if m.get("required") and not mapping.get(f)]
    if missing:
        raise BusinessError(f"Map the required columns: {', '.join(missing)}")
    bad_cols = [c for c in mapping.values() if c and c not in batch.headers]
    if bad_cols:
        raise BusinessError(f"Columns not in file: {', '.join(bad_cols)}")
    batch.mapping = mapping

    uoms = {u.code.upper(): u for u in db.scalars(select(Uom))}
    locations = {l.code.upper(): l for l in db.scalars(select(Location))}
    codes = {c.upper(): i for c, i in db.execute(select(Item.item_code, Item.id))}
    seen_in_file: dict[str, int] = {}
    counts = {"VALID": 0, "INVALID": 0, "DUPLICATE": 0}

    for row in batch.rows:
        errors: list[str] = []
        m = {f: (row.raw.get(col, "") if col else "") for f, col in mapping.items()}
        out: dict = {"name": m.get("name", "").strip()}
        if not out["name"]:
            errors.append("Item name is required")
        raw_type = normalize_name(m.get("item_type", ""))
        itype = TYPE_ALIASES.get(raw_type) or next((t for t in ItemType if normalize_name(t.value) == raw_type), None)
        if not itype:
            errors.append(f"Item type '{m.get('item_type', '')}' is not recognised")
        else:
            out["item_type"] = itype.value
        uom = uoms.get(m.get("uom", "").strip().upper())
        if not uom:
            errors.append(f"Unit '{m.get('uom', '')}' is not configured (known: {', '.join(sorted(uoms))})")
        else:
            out["uom"] = uom.code
        for f in ("min_stock", "reorder_level", "max_stock", "gst_rate", "opening_qty", "unit_cost"):
            d = _dec(m.get(f, ""), f, errors)
            if d is not None:
                out[f] = str(d)
        if uom and out.get("opening_qty"):
            q = Decimal(out["opening_qty"])
            if q != q.quantize(Decimal(1).scaleb(-uom.decimals)):
                errors.append(f"Opening quantity has too many decimals for unit {uom.code}")
        loc_code = m.get("location", "").strip().upper()
        if out.get("opening_qty") and Decimal(out["opening_qty"]) > 0:
            if not loc_code:
                errors.append("Location code is required when opening quantity is given")
            elif loc_code not in locations:
                errors.append(f"Location '{loc_code}' does not exist")
        if loc_code in locations:
            out["location"] = locations[loc_code].code
        for f in ("item_code", "description", "category", "manufacturer", "make_model", "specifications", "hsn_code"):
            if m.get(f):
                out[f] = m[f].strip()

        row.mapped, row.errors, row.duplicate_item_id = out, errors, None
        if errors:
            row.status, row.resolution = "INVALID", None
            counts["INVALID"] += 1
            continue
        dup_id = None
        if out.get("item_code") and out["item_code"].upper() in codes:
            dup_id = codes[out["item_code"].upper()]
            row.errors = [f"Item code {out['item_code']} already exists"]
        else:
            cands = find_duplicates(db, out["name"], out.get("make_model"), out.get("manufacturer"), limit=1)
            if cands:
                dup_id = cands[0].item.id
                row.errors = [f"Possible duplicate of {cands[0].item.item_code} — {cands[0].reason}"]
        key = f"{normalize_name(out['name'])}|{normalize_name(out.get('make_model', ''))}"
        if dup_id is None and key in seen_in_file:
            row.errors = [f"Same item as row {seen_in_file[key]} in this file"]
            row.status, row.resolution = "DUPLICATE", "SKIP"
            counts["DUPLICATE"] += 1
            continue
        seen_in_file.setdefault(key, row.row_no)
        if dup_id:
            row.status, row.duplicate_item_id, row.resolution = "DUPLICATE", dup_id, "LINK"
            counts["DUPLICATE"] += 1
        else:
            row.status, row.resolution = "VALID", "CREATE"
            counts["VALID"] += 1

    batch.status = "VALIDATED"
    batch.summary = {**counts, "total": len(batch.rows)}
    audit(db, user, "import.validate", "import_batch", batch.id, {"mapping": mapping, "summary": batch.summary})
    return batch


def resolve_row(db: Session, user: User, row: ImportRow, resolution: str, link_item_id: int | None) -> ImportRow:
    if row.batch.status != "VALIDATED":
        raise BusinessError("Rows can only be resolved on a validated batch", status_code=409, code="state")
    if row.status == "INVALID":
        raise BusinessError("Fix invalid rows in the source file and re-upload")
    resolution = resolution.upper()
    if resolution not in ("CREATE", "SKIP", "LINK"):
        raise BusinessError("Resolution must be CREATE, SKIP or LINK")
    if resolution == "LINK":
        target = link_item_id or row.duplicate_item_id
        if not target or db.get(Item, target) is None:
            raise BusinessError("Choose an existing item to link to")
        row.duplicate_item_id = target
    if resolution == "CREATE" and row.mapped.get("item_code") and db.scalar(select(Item.id).where(Item.item_code == row.mapped["item_code"])):
        raise BusinessError("This item code already exists — link or skip instead")
    row.resolution = resolution
    audit(db, user, "import.resolve", "import_row", row.id, {"resolution": resolution, "item_id": row.duplicate_item_id})
    return row


def approve_batch(db: Session, user: User, batch: ImportBatch) -> ImportBatch:
    if batch.status != "VALIDATED":
        raise BusinessError("Only validated batches can be approved", status_code=409, code="state")
    if batch.uploaded_by_id == user.id:
        raise BusinessError("Maker–checker: the uploader cannot approve their own import", status_code=403, code="forbidden")
    locations = {l.code: l for l in db.scalars(select(Location))}
    categories = {c.name.lower(): c for c in db.scalars(select(ItemCategory).where(ItemCategory.parent_id.is_(None)))}
    uoms = {u.code: u for u in db.scalars(select(Uom))}
    created = linked = skipped = opening = 0
    for row in batch.rows:
        if row.status == "INVALID" or row.resolution in (None, "SKIP"):
            if row.status != "INVALID":
                row.status = "SKIPPED"
            skipped += 1
            continue
        d = row.mapped
        loc = locations.get(d.get("location", ""))
        if row.resolution == "LINK":
            item = db.get(Item, row.duplicate_item_id)
            linked += 1
        else:
            cat = None
            if d.get("category"):
                cat = categories.get(d["category"].lower())
                if cat is None:
                    cat = ItemCategory(name=d["category"])
                    db.add(cat)
                    db.flush()
                    categories[cat.name.lower()] = cat
            payload = {
                "name": d["name"], "item_type": d["item_type"], "uom_id": uoms[d["uom"]].id,
                "description": d.get("description"), "category_id": cat.id if cat else None,
                "manufacturer": d.get("manufacturer"), "make_model": d.get("make_model"), "specifications": d.get("specifications"),
                "min_stock": Decimal(d.get("min_stock", "0")), "reorder_level": Decimal(d.get("reorder_level", "0")),
                "max_stock": Decimal(d.get("max_stock", "0")), "hsn_code": d.get("hsn_code"),
                "gst_rate": Decimal(d["gst_rate"]) if d.get("gst_rate") else None, "default_location_id": loc.id if loc else None,
            }
            # Duplicates were reviewed row-by-row during validation; the approver confirmed CREATE.
            item = create_item(db, user, payload, reviewed=True)
            if d.get("item_code") and not db.scalar(select(Item.id).where(Item.item_code == d["item_code"], Item.id != item.id)):
                item.item_code = d["item_code"]  # preserve legacy codes from the existing system
            created += 1
        row.result_item_id = item.id
        qty = Decimal(d.get("opening_qty") or 0)
        if qty > 0 and loc:
            post_movement(db, user, MovementInput(
                txn_uuid=f"import-{batch.id}-{row.row_no}", movement_type=MovementType.OPENING, item_id=item.id,
                quantity=qty, to_location_id=loc.id, unit_cost=Decimal(d.get("unit_cost") or 0),
                ref_doc_type="IMPORT", ref_doc_no=batch.number, source="import",
            ), _internal=True)
            opening += 1
        row.status = "IMPORTED"
    batch.status = "APPROVED"
    batch.approved_by_id = user.id
    batch.approved_at = datetime.now(timezone.utc)
    batch.summary = {**batch.summary, "created": created, "linked": linked, "skipped": skipped, "opening_movements": opening}
    audit(db, user, "import.approve", "import_batch", batch.id, batch.summary)
    return batch

