"""Screens carried over from the Oracle Fusion set-up: OSP job-work delivery challans, the production details report,
the inventory overview, cycle counts, item quantities by subinventory, own item numbers and supplier returns."""
from __future__ import annotations

import io
import re
import uuid
from datetime import date

from openpyxl import load_workbook


def ok(r, code=200):
    assert r.status_code == code, r.text
    return r.json()


def _move(api, **body):
    return ok(api.post("/api/stock/movements", json={"txn_uuid": str(uuid.uuid4()), **body}))


def _job(as_user, make_item, refs, qty="20"):
    """A processed spare routed through one in-house step and one outside (job-work) step, with a work order."""
    admin, stores, production = as_user("admin"), as_user("stores"), as_user("production")
    vendor = ok(admin.post("/api/masters/suppliers", json={
        "name": "Sakthi Heat Treatment Industry", "address": "6/115 A, Pannimadai Road, Coimbatore", "gstin": "33ABCDE1234F1Z5"}), 201)
    item = make_item("Tipped Finger DRG.NO. 2303790 - FOR WME", item_type="PROCESSED_SPARE", item_code="M502303790", hsn_code="84483100")["item"]
    ok(stores.put(f"/api/items/{item['id']}/process", json={"steps": [
        {"operation": "Turning", "mode": "IN_HOUSE", "setup_min": "10", "run_min": "2"},
        {"operation": "Hardening and tempering", "mode": "OUTSIDE", "outside_days": 5, "supplier_id": vendor["id"]}]}))
    wo = ok(production.post("/api/work-orders", json={"item_id": item["id"], "quantity": qty}), 201)
    return vendor, item, wo


def test_outside_step_raises_numbered_job_work_challan(as_user, make_item, refs):
    production, qc = as_user("production"), as_user("qc")
    vendor, item, wo = _job(as_user, make_item, refs)
    assert item["item_code"] == "M502303790"  # Oracle item number kept
    ok(production.post(f"/api/work-orders/{wo['id']}/ops/1", json={"action": "start"}))
    ok(production.post(f"/api/work-orders/{wo['id']}/ops/1", json={"action": "finish"}))
    ok(qc.post(f"/api/work-orders/{wo['id']}/ops/1", json={"action": "qc", "passed": True}))

    wo = ok(production.post(f"/api/work-orders/{wo['id']}/ops/2", json={"action": "send", "challan": {
        "po_number": "174637", "po_date": "2026-10-03", "vehicle_no": "tn 66 ak 5769"}}))
    op = wo["ops"][1]
    assert op["status"] == "AT_VENDOR" and re.fullmatch(r"U4-\d{2}-\d{2}/0001", op["challan_no"]) and op["challan_id"]

    register = ok(production.get("/api/osp/challans", params={"po_number": "174637"}))
    assert len(register) == 1 and register[0]["process_nature"] == "HARDENING AND TEMPERING"
    assert register[0]["vehicle_no"] == "TN 66 AK 5769" and register[0]["status"] == "SENT" and float(register[0]["quantity"]) == 20
    ch = ok(production.patch(f"/api/osp/challans/{op['challan_id']}", json={"waybill_no": "EWB 1234 5678 9012", "transporter": "Own vehicle"}))
    assert ch["waybill_no"] == "EWB 1234 5678 9012" and ch["consignee"]["gstin"] == "33ABCDE1234F1Z5"
    assert ch["company"]["gstin"] == "33AAACL3521E1Z6" and ch["lines"][0]["item_code"] == "M502303790" and ch["lines"][0]["hsn_code"] == "84483100"
    assert as_user("sales").patch(f"/api/osp/challans/{op['challan_id']}", json={"vehicle_no": "X"}).status_code == 403

    ok(production.post(f"/api/work-orders/{wo['id']}/ops/2", json={"action": "receive"}))
    assert ok(production.get(f"/api/osp/challans/{op['challan_id']}"))["status"] == "RECEIVED"

    # The unit code on the company profile drives the challan prefix.
    ok(as_user("admin").put("/api/masters/company", json={"unit_code": "U5"}))
    assert ok(production.get("/api/masters/company"))["profile"]["unit_code"] == "U5"
    assert all(not s["key"].startswith("company.") for s in ok(production.get("/api/masters/settings")))


def test_production_details_report_matches_oracle_layout(as_user, make_item, refs):
    production, qc, stores = as_user("production"), as_user("qc"), as_user("stores")
    loc = refs["loc"]["MAIN-R01-B01"]
    vendor, item, wo = _job(as_user, make_item, refs)
    steel = make_item("EN8 Round Bar 25mm", item_type="RAW_MATERIAL")["item"]
    _move(stores, movement_type="RECEIPT", item_id=steel["id"], quantity="50", to_location_id=loc, unit_cost="80")
    ok(stores.post(f"/api/work-orders/{wo['id']}/materials", json={"item_id": steel["id"], "quantity": "10", "location_id": loc}))
    for seq, (go, back) in ((1, ("start", "finish")), (2, ("send", "receive"))):
        ok(production.post(f"/api/work-orders/{wo['id']}/ops/{seq}", json={"action": go, "supplier_id": vendor["id"]}))
        ok(production.post(f"/api/work-orders/{wo['id']}/ops/{seq}", json={"action": back}))
        ok(qc.post(f"/api/work-orders/{wo['id']}/ops/{seq}", json={"action": "qc", "passed": True}))
    ok(production.post(f"/api/work-orders/{wo['id']}/complete", json={"quantity": "5", "location_id": loc, "unit_cost": "60"}))
    ok(production.post(f"/api/work-orders/{wo['id']}/complete", json={"quantity": "15", "location_id": loc, "unit_cost": "60"}))

    today = date.today().isoformat()
    rep = ok(production.get("/api/reports/production", params={"date_from": today, "date_to": today}))
    assert [c["label"] for c in rep["columns"]][:5] == ["Org", "Job Type", "Job No", "Item No", "Item Desc"]
    first, second = rep["rows"]
    assert first["org"] == "U4" and first["job_type"] == "SHOPMADE" and first["item_no"] == "M502303790"
    assert first["remarks"] == "PARTIAL COMPLETION" and second["remarks"] == "COMPLETED"
    assert float(first["transaction_qty"]) == 5 and float(first["start_qty"]) == 20 and float(first["qty_completed"]) == 20
    assert first["job_status"] == "COMPLETED" and float(first["wip_value"]) == 300
    assert float(first["cost_incurred"]) == 800 and float(first["cost_relieved"]) == 1200
    assert rep["totals"] == {**rep["totals"], "transactions": 2, "jobs": 1, "completed": 1}
    assert ok(production.get("/api/reports/production", params={"date_from": today, "date_to": today, "q": "nothing-like-this"}))["rows"] == []

    x = production.get("/api/reports/production.xlsx", params={"date_from": today, "date_to": today})
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["A1"].value == "From Date:" and ws["C4"].value == "Job No" and ws["J5"].value == "PARTIAL COMPLETION"


def test_cycle_count_variance_goes_to_maker_checker(as_user, make_item, refs):
    stores, manager = as_user("stores"), as_user("manager")
    loc = refs["loc"]["MAIN-R01-B01"]
    item = make_item("Ball Bearing 6205")["item"]
    _move(stores, movement_type="RECEIPT", item_id=item["id"], quantity="10", to_location_id=loc, unit_cost="100")

    c = ok(stores.post("/api/inventory/cycle-counts", json={"scope": "r01 b01"}), 201)
    line = next(l for l in c["lines"] if l["item_id"] == item["id"])
    assert line["system_qty"] is None and line["location"] == "MAIN-R01-B01"  # blind count
    assert stores.post(f"/api/inventory/cycle-counts/{c['id']}/submit").status_code == 422  # not counted yet
    ok(stores.put(f"/api/inventory/cycle-counts/{c['id']}/lines", json=[{"line_id": l["id"], "counted_qty": "8" if l["id"] == line["id"] else "0"}
                                                                         for l in c["lines"]]))
    done = ok(stores.post(f"/api/inventory/cycle-counts/{c['id']}/submit"))
    assert done["status"] == "SUBMITTED" and len(done["adjustments"]) == 1
    counted = next(l for l in done["lines"] if l["id"] == line["id"])
    assert float(counted["system_qty"]) == 10 and float(counted["variance"]) == -2

    adj = next(a for a in ok(manager.get("/api/stock/adjustments")) if a["number"] == done["adjustments"][0])
    ok(manager.post(f"/api/stock/adjustments/{adj['id']}/decision", json={"approve": True}))
    assert float(ok(stores.get(f"/api/items/{item['id']}"))["totals"]["on_hand"]) == 8
    assert stores.post("/api/inventory/cycle-counts", json={"scope": "R99"}).status_code == 422


def test_overview_item_quantities_supplier_return_and_item_numbers(as_user, make_item, refs):
    stores = as_user("stores")
    loc = refs["loc"]["MAIN-R01-B01"]
    screw = make_item("CH.HEAD SCREW(Platted) 4 x 10", item_type="RAW_MATERIAL", item_code="A401112106")["item"]
    assert as_user("stores").post("/api/items", json={"name": "Another screw", "item_type": "RAW_MATERIAL", "item_code": "a401112106",
                                                      "uom_id": refs["uom"]["NOS"]}).status_code == 409
    _move(stores, movement_type="RECEIPT", item_id=screw["id"], quantity="1167", to_location_id=loc, unit_cost="1.5")
    _move(stores, movement_type="ISSUE", item_id=screw["id"], quantity="7", from_location_id=loc,
          ref_doc_type="SUPPLIER_RETURN", ref_doc_no="RTV-0001")  # rejected lot back to the supplier

    q = ok(stores.get(f"/api/locator/items/{screw['id']}"))["quantities"]
    assert q["organization"] == "U4" and float(q["on_hand"]) == 1160 and q["uom_name"] == "Numbers"
    assert q["subinventories"][0]["code"] == "MAIN" and q["subinventories"][0]["locators"][0]["code"] == "MAIN-R01-B01"
    assert ok(stores.get("/api/locator/find", params={"query": "A401112106"}))["item_id"] == screw["id"]

    ov = ok(stores.get("/api/inventory/overview"))
    assert ov["organization"]["code"] == "U4"
    assert {"picks", "cycle_counts", "pending", "receipts", "inspection", "shipments", "osp", "restock"} <= set(ov)
    assert as_user("sales").get("/api/inventory/overview").status_code == 200
